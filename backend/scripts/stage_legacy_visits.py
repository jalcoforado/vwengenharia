import argparse
import asyncio
import hashlib
import json
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import get_session_factory
from app.models.identity import Tenant
from app.models.legacy import LegacyVisitStage

LOCAL_TZ = ZoneInfo("America/Fortaleza")


def serialize_value(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def normalize_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def parse_legacy_datetime(value) -> datetime | None:
    if value is None:
        return None

    parsed = None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime(value.year, value.month, value.day)
    elif isinstance(value, str):
        text = value.strip()
        for fmt in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%d/%m/%Y %H:%M:%S",
            "%d/%m/%Y %H:%M",
            "%d/%m/%Y",
        ):
            try:
                parsed = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if parsed is None:
            try:
                parsed = datetime.fromisoformat(text)
            except ValueError:
                return None

    if parsed is None:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=LOCAL_TZ)
    return parsed.astimezone(UTC)


def fingerprint(payload: dict) -> str:
    canonical = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


async def stage_file(tenant_id: str, file_path: Path, sheet_name: str) -> None:
    if not file_path.exists():
        raise RuntimeError(f"file not found: {file_path}")

    workbook = load_workbook(file_path, read_only=True, data_only=False)
    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f"sheet not found: {sheet_name}")

    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [normalize_text(value) or f"COL_{index + 1}" for index, value in enumerate(next(rows))]

    session_factory = get_session_factory()
    async with session_factory() as session:
        tenant = await session.get(Tenant, tenant_id)
        if tenant is None:
            raise RuntimeError("tenant not found")

        existing = set(
            (
                await session.execute(
                    select(LegacyVisitStage.fingerprint).where(
                        LegacyVisitStage.tenant_id == tenant.id
                    )
                )
            ).scalars()
        )

        staged = 0
        skipped = 0
        errors = 0

        for row_number, values in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in values):
                continue

            raw = {
                headers[index]: serialize_value(value)
                for index, value in enumerate(values)
                if index < len(headers)
            }
            row_fingerprint = fingerprint(raw)
            if row_fingerprint in existing:
                skipped += 1
                continue

            station_label = normalize_text(raw.get("Estações"))
            technician_label = normalize_text(raw.get("Técnico"))
            scheduled_at = parse_legacy_datetime(values[6] if len(values) > 6 else None)
            source_status = normalize_text(raw.get("Status"))

            error_code = None
            error_message = None
            stage_status = "STAGED"
            if station_label is None:
                error_code = "MISSING_STATION"
                error_message = "Linha sem rótulo de estação."
            elif technician_label is None:
                error_code = "MISSING_TECHNICIAN"
                error_message = "Linha sem técnico."
            elif scheduled_at is None:
                error_code = "INVALID_VISIT_DATETIME"
                error_message = "Data e Hora ausente ou inválida."

            if error_code:
                stage_status = "ERROR"
                errors += 1

            session.add(
                LegacyVisitStage(
                    tenant_id=tenant.id,
                    source_file=file_path.name,
                    source_row=row_number,
                    fingerprint=row_fingerprint,
                    station_label=station_label,
                    technician_label=technician_label,
                    scheduled_at=scheduled_at,
                    source_status=source_status,
                    raw_payload=raw,
                    status=stage_status,
                    error_code=error_code,
                    error_message=error_message,
                )
            )
            existing.add(row_fingerprint)
            staged += 1

            if staged % 500 == 0:
                await session.commit()

        await session.commit()
        print(
            json.dumps(
                {
                    "tenant_id": str(tenant.id),
                    "source_file": file_path.name,
                    "sheet": sheet_name,
                    "staged": staged,
                    "skipped_existing": skipped,
                    "errors": errors,
                },
                ensure_ascii=False,
            )
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Stage legacy VW Engenharia visits from XLSX without normalizing them."
    )
    parser.add_argument("--tenant-id", required=True)
    parser.add_argument("--file", required=True, type=Path)
    parser.add_argument("--sheet", default="Página1")
    return parser.parse_args()


async def main() -> None:
    args = parse_args()
    await stage_file(args.tenant_id, args.file, args.sheet)


if __name__ == "__main__":
    asyncio.run(main())
